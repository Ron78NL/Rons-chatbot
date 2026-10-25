import os
import json
import tempfile
import re
import zipfile
import io
import base64
from dotenv import load_dotenv
import anthropic
from pypdf import PdfReader
from docx import Document

load_dotenv()

ANTHROPIC_API_KEY = os.getenv("ANTHROPIC_API_KEY")

HISTORY_FILE = "chat_history.json"
SALDO_FILE = "saldo_history.json"

anthropic_client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY) if ANTHROPIC_API_KEY else None

MODELLEN = {
    "Claude Sonnet 5.5 (Slim & Analyse)": {
        "id": "claude-sonnet-5-5",
        "prijs_in": 3.0,
        "prijs_uit": 15.0,
    },
    "Claude Haiku 5.5 (Snel & Goedkoop)": {
        "id": "claude-haiku-5-5",
        "prijs_in": 1.0,
        "prijs_uit": 5.0,
    },
}
STANDAARD_MODEL = "Claude Sonnet 5.5 (Slim & Analyse)"


def bereken_kosten(model_label, input_tokens, output_tokens):
    """Berekent de kosten in USD op basis van het aantal tokens."""
    model = MODELLEN.get(model_label, MODELLEN[STANDAARD_MODEL])
    return (input_tokens * model["prijs_in"] + output_tokens * model["prijs_uit"]) / 1_000_000


def call_ai_model(model_label, messages, system_prompt="", document_context=""):
    """Roept de Anthropic API (zonder streaming) aan met foutafhandeling."""
    if not anthropic_client:
        return "Fout: Geen ANTHROPIC_API_KEY gevonden in je .env bestand of omgevingsvariabelen."

    api_model = MODELLEN.get(model_label, MODELLEN[STANDAARD_MODEL])["id"]

    try:
        formatted_messages = []
        for m in messages:
            role = m.get("role", "user")
            if role not in ("user", "assistant"):
                role = "user"
            formatted_messages.append({"role": role, "content": m.get("content", "")})

        system_instruction = system_prompt
        if document_context:
            system_instruction += f"\n\nDocument Context:\n{document_context}"

        kwargs = {
            "model": api_model,
            "max_tokens": 8000,
            "messages": formatted_messages,
        }
        if system_instruction.strip():
            kwargs["system"] = system_instruction.strip()

        response = anthropic_client.messages.create(**kwargs)
        return response.content[0].text
    except Exception as e:
        return f"Anthropic API Fout opgetreden: {str(e)}"


def extraheer_tekst(bestand):
    """Extraheert tekst uit geüploade PDF, TXT, DOCX of ZIP bestanden."""
    if not bestand:
        return ""

    toegestane_extensies = (".txt", ".pdf", ".docx", ".zip", ".png", ".jpg", ".jpeg", ".webp")
    naam = getattr(bestand, "name", "onbekend")
    bestandsnaam = naam.lower()

    if not bestandsnaam.endswith(toegestane_extensies):
        raise ValueError(f"Bestandstype niet toegestaan: {naam}")

    try:
        if bestandsnaam.endswith(".txt"):
            content = bestand.getvalue()
            if isinstance(content, bytes):
                return content.decode("utf-8", errors="ignore")
            return str(content)

        if bestandsnaam.endswith(".pdf"):
            reader = PdfReader(bestand)
            tekst = ""
            for pagina in reader.pages:
                t = pagina.extract_text()
                if t:
                    tekst += t + "\n"
            if not tekst.strip():
                raise ValueError("De PDF bevat geen leesbare tekst (mogelijk een gescande afbeelding).")
            return tekst.strip()

        if bestandsnaam.endswith(".docx"):
            doc = Document(bestand)
            delen = [p.text for p in doc.paragraphs if p.text.strip()]
            for tabel in doc.tables:
                for rij in tabel.rows:
                    cellen = [c.text.strip() for c in rij.cells]
                    if any(cellen):
                        delen.append(" | ".join(cellen))
            tekst = "\n".join(delen).strip()
            if not tekst:
                raise ValueError("Het Word-bestand bevat geen leesbare tekst.")
            return tekst

        if bestandsnaam.endswith(".zip"):
            gepakte_inhoud = ""
            with zipfile.ZipFile(io.BytesIO(bestand.getvalue())) as z:
                for bestands_info in z.namelist():
                    if bestands_info.endswith('/') or '__MACOSX' in bestands_info or os.path.basename(bestands_info).startswith('.'):
                        continue
                    if bestands_info.lower().endswith(('.txt', '.py', '.md', '.json', '.html', '.css', '.js', '.csv')):
                        try:
                            with z.open(bestands_info) as f:
                                file_content = f.read().decode('utf-8', errors='ignore')
                                gepakte_inhoud += f"\n\n--- Bestand in ZIP: {bestands_info} ---\n{file_content}"
                        except Exception:
                            pass
            if not gepakte_inhoud.strip():
                raise ValueError("Geen leesbare tekst- of codebestanden gevonden in het ZIP-archief.")
            return gepakte_inhoud.strip()

    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Kon de inhoud van {naam} niet uitlezen: {str(e)}")

    return ""


def verwerk_afbeelding(bestand):
    """Leest een geüploade afbeelding uit en zet deze om naar base64 voor de Anthropic API."""
    if not bestand:
        return None
    
    naam = getattr(bestand, "name", "onbekend").lower()
    media_type = "image/jpeg"
    if naam.endswith(".png"):
        media_type = "image/png"
    elif naam.endswith(".webp"):
        media_type = "image/webp"
    elif naam.endswith(".gif"):
        media_type = "image/gif"
        
    try:
        binaire_data = bestand.getvalue()
        base64_data = base64.b64encode(binaire_data).decode("utf-8")
        return {
            "type": "image",
            "source": {
                "type": "base64",
                "media_type": media_type,
                "data": base64_data,
            }
        }
    except Exception as e:
        raise ValueError(f"Kon afbeelding {naam} niet verwerken: {str(e)}")


def safe_save_json(filepath, data):
    dir_name = os.path.dirname(filepath)
    if dir_name and not os.path.exists(dir_name):
        os.makedirs(dir_name, exist_ok=True)

    fd, temp_path = tempfile.mkstemp(dir=dir_name if dir_name else ".")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=4)
        os.replace(temp_path, filepath)
        return True
    except Exception:
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except Exception:
                pass
        return False


def safe_load_json(filepath, default_value=None):
    if default_value is None:
        default_value = {}
    if not os.path.exists(filepath):
        return default_value
    try:
        with open(filepath, "r", encoding="utf-8") as f:
            content = f.read().strip()
            if not content:
                return default_value
            return json.loads(content)
    except Exception:
        return default_value


def extracteer_codeblokken(tekst):
    """Scant de AI-respons op codeblokken en herkent optionele bestandsnamen."""
    if not tekst:
        return []

    patroon = r"```([\w+#-]*)[ \t]*\n(.*?)```"
    matches = re.findall(patroon, tekst, re.DOTALL)
    code_blokken = []
    
    for idx, (taal, code) in enumerate(matches):
        bestandsnaam = f"bestand_{idx+1}.{taal.lower() if taal else 'txt'}"
        code_inhoud = code.strip()
        
        Eerste_regel = code_inhoud.split('\n')[0] if code_inhoud else ""
        if "filename:" in Eerste_regel.lower() or "bestand:" in Eerste_regel.lower():
            delen = re.split(r"[:=]", Eerste_regel, 1)
            if len(delen) > 1:
                bestandsnaam = delen[1].strip().strip("'\"#//")
                code_inhoud = "\n".join(code_inhoud.split('\n')[1:]).strip()

        code_blokken.append({
            "taal": (taal or "").lower(),
            "code": code_inhoud,
            "naam": bestandsnaam
        })

    if not code_blokken and "<html" in tekst.lower():
        code_blokken.append({
            "taal": "html",
            "code": tekst.strip(),
            "naam": "index.html"
        })

    return code_blokken
